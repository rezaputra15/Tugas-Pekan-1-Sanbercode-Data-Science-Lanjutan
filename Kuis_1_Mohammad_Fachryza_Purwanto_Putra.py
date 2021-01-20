from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
import csv, pandas as pd

class kuis:
    def __init__(self, data_series, raw_data, hasil_data):
        self.data_series = data_series
        self.raw_data = raw_data
        self.hasil_data = hasil_data

    def clean (self):
        df_1 = pd.read_csv(self.raw_data)
        df_2 = pd.read_csv(self.raw_data)

        df_1['tags'] = df_1['Country Name'].apply(lambda x: x.lower().replace(" ",""))
        df_2['tags'] = df_2['Country Name'].apply(lambda x: x.lower().replace(" ",""))

        gdp = df_1[['Country Name', 'Indicator Code', '2010', 'tags']].copy()
        gdp.rename(columns = {'Indicator Code': 'Indicator_Code'}, inplace = True)
        gdp = gdp[gdp.Indicator_Code == 'NA.GDP.EXC.OG.CR']
        gdp = gdp.fillna(0)
        #print(gdp.head())

        hdi = df_2[['Indicator Code', '2010', 'tags']].copy()
        hdi.rename(columns = {'Indicator Code': 'Indicator_Code'}, inplace = True)
        hdi = hdi[hdi.Indicator_Code == 'IDX.HDI']
        hdi = hdi.fillna(0)
        #print(hdi.head())

        df = pd.merge(gdp, hdi, on='tags')
        #print(df.head())

        df_clean = df.drop(columns=['Indicator_Code_x', 'tags', 'Indicator_Code_y'])
        df_clean.rename(columns = {'Country Name': 'Nama Provinsi', '2010_x':'GDP', '2010_y':'HDI'}, inplace = True)
        #print(df_clean.head())

        df_clean1 = df_clean[['GDP']].copy()
        df_norm = (df_clean1 - df_clean1.min()) / (df_clean1.max() - df_clean1.min())
        df_norm.rename(columns = {'GDP': 'GDP Normalisasi'}, inplace = True)
        #print(df_norm.head())

        new_df = df_clean.join(df_norm)
        new_df = new_df.drop(columns=['GDP'])
        #print(new_df.head())

        new_df.to_excel(self.hasil_data, index = False)

    def visualisasi(self):
        wb = load_workbook(filename=self.hasil_data)
        ws = wb.active
        chart1 = BarChart()
        chart1.type = "col"
        chart1.style = 5
        chart1.title = "GDP Provinsi"
        chart1.y_axis.title = "Total GDP"
        chart1.x_axis.title = "Provinsi"

        data = Reference(ws, min_col=3, min_row=1, max_row=35, max_col=3)
        cats = Reference(ws, min_col=1, min_row=2, max_row=35)

        chart1.height = 10 #default is 7.5
        chart1.width = 30 #default is 15
        chart1.add_data(data, titles_from_data=True)
        chart1.set_categories(cats)
        
        chart2 = BarChart()
        chart2.type = "col"
        chart2.style = 3
        chart2.title = "HDI Provinsi"
        chart2.y_axis.title = "HDI"
        chart2.x_axis.title = "Provinsi"

        data = Reference(ws, min_col=2, min_row=1, max_row=35, max_col=2)
        cats = Reference(ws, min_col=1, min_row=2, max_row=35)

        chart2.height = 10 #default is 7.5
        chart2.width = 30 #default is 15
        chart2.add_data(data, titles_from_data=True)
        chart2.set_categories(cats)
        
        ws.add_chart(chart1, "E2")
        ws.add_chart(chart2, "E25")
        wb.save(hasil_data)

data_series = r"D:\Projectku\Sanbercode-Data_Science_Lanjutan\Pekan_1\Bahan\Kuis_1\INDODAPOERSeries.csv"
raw_data = r"D:\Projectku\Sanbercode-Data_Science_Lanjutan\Pekan_1\Bahan\Kuis_1\raw_data.csv"
hasil_data = r"D:\Projectku\Sanbercode-Data_Science_Lanjutan\Pekan_1\Bahan\Kuis_1\fachreyzaputra.xlsx"

hasil = kuis(data_series, raw_data, hasil_data)
hasil.clean()
hasil.visualisasi()